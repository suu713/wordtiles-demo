import argparse
import hashlib
import json
import re
from pathlib import Path

from openpyxl import load_workbook


REQUIRED_HEADERS = ["单词", "范围库", "音标", "是否缩写", "英文释义", "中文释义"]
SHEETS = {"必答词": "required", "额外词": "extra"}


def read_sheet(sheet, role):
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value).strip() if value is not None else "" for value in next(rows, ())]
    if len(set(headers)) != len(headers):
        raise ValueError(f"工作表 {sheet.title} 的表头包含重复列")
    missing = set(REQUIRED_HEADERS) - set(headers)
    if missing or not {"分类", "细分类"}.intersection(headers) or not {"生僻程度(1-5)", "生僻档位"}.intersection(headers):
        raise ValueError(f"工作表 {sheet.title} 缺少必需列：{sorted(missing)}，且须有分类与生僻档位")
    category_header = "细分类" if "细分类" in headers else "分类"
    entries = []
    seen = set()
    for line, values in enumerate(rows, start=2):
        row = dict(zip(headers, values))
        raw_word = row["单词"]
        if raw_word is None or (isinstance(raw_word, str) and not raw_word.strip()):
            continue
        # Excel can store the literal English headword FALSE as a boolean cell.
        word = ("TRUE" if raw_word is True else "FALSE" if raw_word is False else str(raw_word).strip()).upper()
        if not re.fullmatch(r"[A-Z]{3,8}", word):
            raise ValueError(f"{sheet.title} 第{line}行: 非3至8字母英文词 {word!r}")
        if word in seen:
            raise ValueError(f"{sheet.title} 第{line}行: 重复词 {word}")
        seen.add(word)
        range_label = str(row["范围库"] or "").strip()
        range_match = re.fullmatch(r"([1-4])级", range_label)
        if not range_match:
            raise ValueError(f"{sheet.title} 第{line}行: 范围库应为1至4级")
        rarity_raw = row.get("生僻档位", row.get("生僻程度(1-5)"))
        rarity_match = re.fullmatch(r"([1-5])档", str(rarity_raw or "").strip())
        rarity = int(rarity_match.group(1)) if rarity_match else rarity_raw
        if type(rarity) is not int or rarity not in range(1, 6):
            raise ValueError(f"{sheet.title} 第{line}行: 生僻档位应为1至5档")
        abbreviation = str(row["是否缩写"] or "").strip()
        if abbreviation not in ("是", "否"):
            raise ValueError(f"{sheet.title} 第{line}行: 是否缩写应为是或否")
        fields = ["音标", category_header, "英文释义", "中文释义"]
        if any(not str(row[name] or "").strip() for name in fields):
            raise ValueError(f"{sheet.title}!{line}: 词条注释字段缺失")
        entries.append({
            "word": word,
            "role": role,
            "range": range_label,
            "rangeTier": int(range_match.group(1)),
            "phonetic": str(row["音标"]).strip(),
            "category": str(row[category_header]).strip(),
            "abbreviation": abbreviation == "是",
            "rarity": rarity,
            "definitionEn": str(row["英文释义"]).strip(),
            "definitionZh": str(row["中文释义"]).strip(),
            **{key: str(row.get(header) or "").strip() for key, header in {
                "emotion": "情绪类型", "partOfSpeech": "词性", "rarityLabel": "生僻档位",
                "rarityDescription": "档位说明", "rarityStatus": "档位状态",
                "abbreviationType": "缩写类型", "abbreviationExpansion": "缩写全称",
                "reviewNote": "复核提示", "proofreadingNote": "本次校对记录",
            }.items()},
            "confidence": {key: str(row.get(header) or "").strip() for key, header in {
                "category": "分类置信度", "partOfSpeech": "词性置信度",
                "abbreviation": "缩写置信度", "rarity": "档位置信度",
            }.items()},
        })
    return entries


def main():
    parser = argparse.ArgumentParser(description="将 WT 词库工作簿打包为网页运行时 JSON")
    parser.add_argument("input", type=Path, help="wt词库.xlsx")
    parser.add_argument("--output", type=Path, default=Path("dist/wt-word-bank.json"))
    args = parser.parse_args()
    source_bytes = args.input.read_bytes()
    workbook = load_workbook(args.input, read_only=True, data_only=True)
    missing = set(SHEETS) - set(workbook.sheetnames)
    if missing:
        raise ValueError("缺少工作表：" + ", ".join(sorted(missing)))
    bank = {role: read_sheet(workbook[name], role) for name, role in SHEETS.items()}
    workbook.close()
    required = {entry["word"] for entry in bank["required"]}
    extras = {entry["word"] for entry in bank["extra"]}
    overlap = required & extras
    if overlap:
        raise ValueError("必答词与额外词重复：" + ", ".join(sorted(overlap)[:20]))
    digest = hashlib.sha256(source_bytes).hexdigest()
    payload = {
        "schemaVersion": 2,
        "version": f"wt-workbook-{digest[:12]}",
        "language": "english",
        "source": {
            "file": args.input.name,
            "sha256": digest,
            "requiredSheet": "必答词",
            "extraSheet": "额外词",
            "requiredCount": len(bank["required"]),
            "extraCount": len(bank["extra"]),
        },
        **bank,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    staged = args.output.with_suffix(args.output.suffix + ".tmp")
    staged.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")
    staged.replace(args.output)
    print(f"写入 {args.output}: {len(required):,} 必答词 + {len(extras):,} 额外词 · SHA256 {digest}")


if __name__ == "__main__":
    main()
